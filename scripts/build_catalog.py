"""Build the video catalog: one row per generation input, grouped by data source.

Usage: .venv/bin/python scripts/build_catalog.py

Writes metamanifest/video_catalog.xlsx (colour-banded sections per data source, gate results
coloured, frozen header, links to the viewer and the repository) and metamanifest/video_catalog.csv
(the same rows, plain, for scripts). Run scripts/quality_gate.py first.
"""
import csv
import json
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from build_site import DATASETS, ROOT, SITE_LIMIT, collect, gh, probe

VIEWER = "https://wilsonnijc-bot.github.io/robot_object_deformation_detect/robotvideo/"
OUT = ROOT / "metamanifest"
# Section colours per data source: (band, row tint).
COLOURS = {
    "NVIDIA_GR1_100": ("2F6B4F", "E8F2EC"),
    "Humanoid_Everyday": ("8A4B2A", "F6ECE5"),
    "TAVP": ("3B5B8C", "E9EEF6"),
    "ActionNet": ("7A5A9A", "F0EBF5"),
    "RoboMIND": ("A0782A", "F7F0E1"),
    "RoboCOIN_Airbot_MMK2_stack_block": ("2E7A80", "E5F2F3"),
}
GATE_FILL = {"pass": "CDEBD5", "exception": "F9E3B4", "fail": "F5C9C2"}
COLUMNS = [  # (header, width)
    ("#", 5), ("Data source", 18), ("Episode folder", 34), ("Robot", 26), ("Camera view", 18),
    ("Resolution", 11), ("Frames", 8), ("FPS", 6), ("Duration (s)", 10), ("Clip", 16),
    ("Gate", 10), ("Sharpness", 10), ("Threshold", 9),
    ("Action generation prompt", 70), ("Full-sequence prompt", 50), ("Publisher annotation", 45),
    ("Observed outcome", 50), ("Prompt protocol", 18), ("Viewer page", 14), ("Repository folder", 14),
]


def outcome(review):
    for key in ("observed_outcome", "selected_action_outcome", "outcome", "observed_terminal_state", "terminal_state"):
        value = review.get(key)
        if isinstance(value, dict):
            return "; ".join(str(v) for v in (value.get("status"), value.get("description")) if v)
        if value:
            return str(value)
    return review.get("selected_execution_status", "")


def rows():
    shown = {}
    for e in collect():
        n = shown.setdefault(e["dataset"], 0)
        shown[e["dataset"]] = n + 1
        review = json.loads(e["review"]) if e["review"] else {}
        info = probe(e["video"])
        meta, gate = e["meta"], e["gate"]
        on_site = SITE_LIMIT.get(e["dataset"]) is None or n < SITE_LIMIT[e["dataset"]]
        yield {
            "Data source": DATASETS[e["dataset"]]["title"], "dataset": e["dataset"], "Episode folder": e["folder"],
            "Robot": DATASETS[e["dataset"]]["robot"],
            "Camera view": gate["view"].replace("_", "-") + (f" · {meta['view']}" if isinstance(meta.get("view"), str)
                                                            and meta["view"].startswith(("camera", "observation")) else ""),
            "Resolution": f"{gate['width']}×{gate['height']}", "Frames": info["frames"], "FPS": info["fps"],
            "Duration (s)": info["duration"], "Clip": (meta.get("variant") or "").replace("_", " "),
            "Gate": gate["gate"], "Sharpness": float(gate["sharpness"]), "Threshold": int(gate["threshold"]),
            "Action generation prompt": e["prompt"] or "", "Full-sequence prompt": e["full_prompt"] or "",
            "Publisher annotation": e["publisher"] or "", "Observed outcome": outcome(review),
            "Prompt protocol": review.get("prompt_protocol", ""),
            "Viewer page": f"{VIEWER}{e['dataset']}/{e['folder']}/index.html" if on_site else "",
            "Repository folder": gh("tree", e["ep_root"].relative_to(ROOT)),
        }


def main():
    order = list(DATASETS)
    data = sorted(rows(), key=lambda r: order.index(r["dataset"]))
    headers = [h for h, _ in COLUMNS]
    with open(OUT / "video_catalog.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["dataset"] + headers[1:], extrasaction="ignore")
        w.writeheader()
        w.writerows(data)

    wb = Workbook()
    ws = wb.active
    ws.title = "Videos"
    thin = Side(style="thin", color="D9D4CA")
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2B2A27")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    number = 0
    for ds in DATASETS:
        items = [r for r in data if r["dataset"] == ds]
        if not items:
            continue
        band, tint = COLOURS[ds]
        meta, gates = DATASETS[ds], Counter(r["Gate"] for r in items)
        ws.append([f"{meta['title']}  ·  {len(items)} videos  ·  {meta['robot']}  ·  {meta['view']}  ·  "
                   f"{', '.join(f'{v} {k}' for k, v in gates.items())}  ·  {meta['license']}"])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
        ws.cell(r, 1).font = Font(bold=True, color="FFFFFF", size=12)
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=band)
        ws.cell(r, 1).alignment = Alignment(vertical="center")
        ws.row_dimensions[r].height = 24
        for item in items:
            number += 1
            values = [number] + [item[h] for h in headers[1:]]
            ws.append(values)
            r = ws.max_row
            for c in ws[r]:
                c.fill = PatternFill("solid", fgColor=tint)
                c.alignment = wrap
                c.border = Border(bottom=thin)
            gate = ws.cell(r, headers.index("Gate") + 1)
            gate.fill = PatternFill("solid", fgColor=GATE_FILL[item["Gate"]])
            gate.font = Font(bold=True)
            for col, label in (("Viewer page", "open viewer"), ("Repository folder", "open folder")):
                cell = ws.cell(r, headers.index(col) + 1)
                if cell.value:
                    cell.hyperlink, cell.value = cell.value, label
                    cell.font = Font(color="1F5C8B", underline="single")
            prompt = ws.cell(r, headers.index("Action generation prompt") + 1)
            if not prompt.value:
                prompt.value, prompt.font = "(no reviewed prompt yet)", Font(italic=True, color="8A857C")
    for i, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "D2"

    legend = wb.create_sheet("Legend")
    legend.append(["Data source", "Videos", "Gate pass", "Gate exception", "Gate fail", "License", "Source"])
    for c in legend[1]:
        c.font = Font(bold=True)
    for ds in DATASETS:
        items = [r for r in data if r["dataset"] == ds]
        if items:
            g = Counter(r["Gate"] for r in items)
            legend.append([DATASETS[ds]["title"], len(items), g["pass"], g["exception"], g["fail"],
                           DATASETS[ds]["license"], DATASETS[ds]["source"]])
            legend.cell(legend.max_row, 1).fill = PatternFill("solid", fgColor=COLOURS[ds][0])
            legend.cell(legend.max_row, 1).font = Font(bold=True, color="FFFFFF")
    legend.append([])
    for g, fill in GATE_FILL.items():
        legend.append([f"Gate: {g}"])
        legend.cell(legend.max_row, 1).fill = PatternFill("solid", fgColor=fill)
    legend.append(["Sharpness: Laplacian variance of the upscaled conditioning frame; threshold 120 third-person, "
                   "60 first-person. See data_quality_gate.md."])
    for col, width in zip("ABCDEFG", (34, 8, 10, 14, 10, 30, 60)):
        legend.column_dimensions[col].width = width
    wb.save(OUT / "video_catalog.xlsx")
    print(f"{len(data)} videos -> {OUT / 'video_catalog.xlsx'} and video_catalog.csv")


if __name__ == "__main__":
    main()
